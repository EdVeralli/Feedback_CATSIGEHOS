"""
Feedback CATSIGEHOS - Encuesta de satisfaccion de la CAT de turnos SIGEHOS (Boti)

Calcula, por mes y para el total del periodo, las metricas de la encuesta
post-atencion de CATs filtrando las sesiones cuyo ultimo tema para NPS
(variable `ultimotemanps`) es el tema indicado (default: CATSIGEHOS):

    Total_Usuarios   = sesiones que respondieron Si/No (pregunta de efectividad)
    TasaEfectividad  = Si / (Si + No)
    Esfuerzo         = promedio ponderado Muy facil=1, Facil=2, Dificil=4,
                       Muy dificil=5. Se EXCLUYE "Ni facil ni dificil" para
                       replicar el calculo de referencia (ver README).
    Satisfaccion %   = (Muy conforme + Conforme) /
                       (Muy conforme + Conforme + Inconforme + Muy inconforme + No se)

Ademas exporta el texto libre de sugerencias/comentarios.

Uso:
    cd C:\\GCBA\\Metricas_Boti_Mensual\\Feedback_CATSIGEHOS
    python Feedback_CATSIGEHOS.py --anio 2026 --meses 7 8 9
    python Feedback_CATSIGEHOS.py --anio 2026 --meses 8 --tema CATSoporteMiBA

Requiere sesion AWS activa con rol PIBADataScientist:
    aws-azure-login --profile default --mode=gui
"""

import argparse
import os
import sys
import unicodedata
from calendar import monthrange
from datetime import datetime

import awswrangler as wr
import boto3
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ==================== CONFIGURACION ====================
CONFIG = {
    'region': 'us-east-1',
    'workgroup': 'Production-caba-piba-athena-boti-group',
    'database': 'caba-piba-consume-zone-db',
}

TABLA_MSG = '"caba-piba-consume-zone-db"."boti_message_metrics_2"'
TABLA_VARS = '"caba-piba-consume-zone-db"."boti_user_vars_metrics_2"'

VAR_TEMA = 'ultimotemanps'
VARS_TEXTO = ['feedbacksugerencia', 'feedbackcomentario']

# Reglas de la encuesta de CATs en Botmaker (match exacto, normalizado con strip)
REGLAS = {
    # Efectividad (P2)
    'si':             'CXF01CUX04 Sí CATs',
    'no':             'CXF01CUX04 No CATs',
    # Satisfaccion (P3)
    'muy_conforme':   'CXF01CUX04 Muy conforme CATs',
    'conforme':       'CXF01CUX04 Conforme CATs',
    'inconforme':     'CXF01CUX04 Inconforme CATs',
    'muy_inconforme': 'CXF01CUX04 Muy inconforme CATs',
    'no_se':          'CXF01CUX04 No sé CATs',
    # Esfuerzo (P4)
    'muy_facil':      'CXF01CUX04 Muy fácil CATs',
    'facil':          'CXF01CUX04 Fácil CATs',
    'neutro':         'CXF01CUX04 Ni fácil ni difícil CATs',
    'dificil':        'CXF01CUX04 Difícil CATs',
    'muy_dificil':    'CXF01CUX04 Muy difícil CATs',
}

# Pesos del Esfuerzo. 'neutro' NO entra en el calculo (replica la referencia).
PESOS_ESFUERZO = {'muy_facil': 1, 'facil': 2, 'dificil': 4, 'muy_dificil': 5}

MESES_ES = ['', 'ENERO', 'FEBRERO', 'MARZO', 'ABRIL', 'MAYO', 'JUNIO', 'JULIO',
            'AGOSTO', 'SEPTIEMBRE', 'OCTUBRE', 'NOVIEMBRE', 'DICIEMBRE']


# ==================== UTILIDADES ====================
def norm(s):
    """Normaliza un rule_name para comparar: strip + NFC (tildes consistentes)."""
    return unicodedata.normalize('NFC', str(s)).strip()


def log(msg):
    print('[{}] {}'.format(datetime.now().strftime('%H:%M:%S'), msg), flush=True)


def correr_query(sql, session):
    try:
        return wr.athena.read_sql_query(
            sql=sql, database=CONFIG['database'], workgroup=CONFIG['workgroup'],
            boto3_session=session, ctas_approach=False, unload_approach=False)
    except Exception as e:
        if 'workgroup' in str(e).lower():
            log('[ADVERTENCIA] Reintentando sin workgroup...')
            return wr.athena.read_sql_query(
                sql=sql, database=CONFIG['database'], boto3_session=session,
                ctas_approach=False, unload_approach=False)
        raise


def rango_fechas(anio, meses):
    ini = '{:04d}-{:02d}-01'.format(anio, min(meses))
    fin = '{:04d}-{:02d}-{:02d}'.format(anio, max(meses), monthrange(anio, max(meses))[1])
    return ini, fin


# ==================== QUERIES ====================
def cte_sesiones(anio, meses, tema):
    meses_sql = ', '.join("'{}'".format(m) for m in meses)
    return """ses AS (
  SELECT DISTINCT session_id
  FROM {tv}
  WHERE year = '{anio}' AND month IN ({meses})
    AND vars_name = '{var}' AND vars_value = '{tema}'
)""".format(tv=TABLA_VARS, anio=anio, meses=meses_sql, var=VAR_TEMA, tema=tema)


def query_reglas(anio, meses, tema):
    ini, fin = rango_fechas(anio, meses)
    return """WITH {cte}
SELECT month(CAST(m.session_creation_time AS DATE)) AS mes,
       m.rule_name,
       count(distinct m.session_id) AS sesiones,
       count(distinct SUBSTR(m.session_id, 1, 20)) AS usuarios
FROM {tm} m
JOIN ses ON ses.session_id = m.session_id
WHERE CAST(m.session_creation_time AS DATE) BETWEEN date '{ini}' AND date '{fin}'
  AND m.rule_name LIKE 'CXF01CUX04%'
GROUP BY 1, 2
ORDER BY 1, 2""".format(cte=cte_sesiones(anio, meses, tema), tm=TABLA_MSG, ini=ini, fin=fin)


def query_textos(anio, meses, tema):
    meses_sql = ', '.join("'{}'".format(m) for m in meses)
    vars_sql = ', '.join("'{}'".format(v) for v in VARS_TEXTO)
    return """WITH {cte}
SELECT DISTINCT
       v.month AS mes,
       v.session_id,
       CONCAT('https://go.botmaker.com/#/chats/', SUBSTR(v.session_id, 1, 20)) AS link,
       v.vars_name AS tipo,
       v.vars_value AS texto
FROM {tv} v
JOIN ses ON ses.session_id = v.session_id
WHERE v.year = '{anio}' AND v.month IN ({meses})
  AND v.vars_name IN ({vars})
  AND v.vars_value IS NOT NULL AND trim(v.vars_value) <> ''
ORDER BY 1, 2""".format(cte=cte_sesiones(anio, meses, tema), tv=TABLA_VARS,
                         anio=anio, meses=meses_sql, vars=vars_sql)


# ==================== CALCULO ====================
def conteos(df_mes):
    """Devuelve {clave: sesiones} para las reglas esperadas + lista de faltantes."""
    mapa = {norm(r): s for r, s in zip(df_mes['rule_name'], df_mes['sesiones'])}
    c, faltantes = {}, []
    for clave, regla in REGLAS.items():
        if norm(regla) in mapa:
            c[clave] = int(mapa[norm(regla)])
        else:
            c[clave] = 0
            faltantes.append(regla)
    return c, faltantes


def metricas(c):
    tot_ef = c['si'] + c['no']
    pos = c['muy_conforme'] + c['conforme']
    tot_sat = pos + c['inconforme'] + c['muy_inconforme'] + c['no_se']
    tot_esf = sum(c[k] for k in PESOS_ESFUERZO)
    suma_esf = sum(c[k] * p for k, p in PESOS_ESFUERZO.items())
    return {
        'Total_Usuarios': tot_ef,
        'TasaEfectividad': c['si'] / tot_ef if tot_ef else None,
        'Esfuerzo': suma_esf / tot_esf if tot_esf else None,
        'Satisfaccion': pos / tot_sat if tot_sat else None,
    }


# ==================== EXCEL ====================
GRIS = PatternFill('solid', fgColor='D9D9D9')
BORDE = Border(*(Side(style='thin'),) * 4)
NEGRITA_IT = Font(bold=True, italic=True)


def bloque_resumen(ws, fila, titulo, m):
    ws.cell(fila, 1, titulo).font = NEGRITA_IT
    ws.cell(fila, 1).alignment = Alignment(horizontal='center')
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=4)
    hdr = ['Total_Usuarios', 'TasaEfectividad', 'Esfuerzo', 'Satisfacción %']
    for j, h in enumerate(hdr, 1):
        cel = ws.cell(fila + 1, j, h)
        cel.font, cel.fill, cel.border = NEGRITA_IT, GRIS, BORDE
    vals = [m['Total_Usuarios'], m['TasaEfectividad'], m['Esfuerzo'], m['Satisfaccion']]
    fmts = ['#,##0', '0.00%', '0.00', '0.00%']
    for j, (v, f) in enumerate(zip(vals, fmts), 1):
        cel = ws.cell(fila + 2, j, v)
        cel.number_format, cel.border = f, BORDE
    return fila + 4


def escribir_excel(path, tema, resumen, df_det, df_txt):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Resumen'
    fila = 1
    ws.cell(fila, 1, 'Feedback {}'.format(tema)).font = Font(bold=True, size=13)
    fila += 2
    for titulo, m in resumen:
        fila = bloque_resumen(ws, fila, titulo, m)
    ws.cell(fila, 1, 'Esfuerzo: escala 1 (muy fácil) a 5 (muy difícil), sin "Ni fácil ni difícil".').font = Font(italic=True, size=9)
    for j in range(1, 5):
        ws.column_dimensions[get_column_letter(j)].width = 20

    for nombre, df in [('Detalle reglas', df_det), ('Sugerencias', df_txt)]:
        wsx = wb.create_sheet(nombre)
        if df is None or df.empty:
            wsx.cell(1, 1, 'Sin datos')
            continue
        for j, col in enumerate(df.columns, 1):
            cel = wsx.cell(1, j, col)
            cel.font, cel.fill = Font(bold=True), GRIS
            wsx.column_dimensions[get_column_letter(j)].width = 18
        for i, row in enumerate(df.itertuples(index=False), 2):
            for j, v in enumerate(row, 1):
                wsx.cell(i, j, v)
        if nombre == 'Sugerencias':
            wsx.column_dimensions[get_column_letter(len(df.columns))].width = 90
        wsx.freeze_panes = 'A2'
    wb.save(path)


# ==================== MAIN ====================
def main():
    ap = argparse.ArgumentParser(description='Feedback de la encuesta CATs filtrado por ultimotemanps.')
    ap.add_argument('--anio', type=int, required=True)
    ap.add_argument('--meses', type=int, nargs='+', required=True, help='Ej: 7 8 9 (mismo año)')
    ap.add_argument('--tema', default='CATSIGEHOS', help='Valor de ultimotemanps (default CATSIGEHOS)')
    ap.add_argument('--sin-sugerencias', action='store_true', help='No baja el texto libre')
    args = ap.parse_args()

    meses = sorted(set(args.meses))
    if any(m < 1 or m > 12 for m in meses):
        sys.exit('Meses fuera de rango.')

    log('FEEDBACK {} | {} meses {}'.format(args.tema, args.anio, meses))
    session = boto3.Session(region_name=CONFIG['region'])

    log('Query reglas CXF01CUX04 ...')
    df = correr_query(query_reglas(args.anio, meses, args.tema), session)
    if df.empty:
        sys.exit('[ERROR] La query no devolvió reglas. Revisar tema/período.')
    df['mes'] = df['mes'].astype(int)

    resumen_mes, total, hay_faltantes = [], {k: 0 for k in REGLAS}, False
    for mes in meses:
        c, falt = conteos(df[df['mes'] == mes])
        for regla in falt:
            hay_faltantes = True
            log('    [ADVERTENCIA] {} - no se encontró la regla: {}'.format(MESES_ES[mes], regla))
        for k in total:
            total[k] += c[k]
        resumen_mes.append((MESES_ES[mes], metricas(c)))

    titulo_total = 'Últimos {} meses'.format(len(meses)) if len(meses) > 1 else None
    resumen = ([(titulo_total, metricas(total))] if titulo_total else []) + resumen_mes

    df_txt = None
    if not args.sin_sugerencias:
        log('Query sugerencias/comentarios ...')
        df_txt = correr_query(query_textos(args.anio, meses, args.tema), session)

    out_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'output')
    os.makedirs(out_dir, exist_ok=True)
    sufijo = '{}_{}_{}'.format(args.tema, args.anio, '-'.join(str(m) for m in meses))
    path_xlsx = os.path.join(out_dir, 'feedback_{}.xlsx'.format(sufijo))
    df_det = df.sort_values(['mes', 'rule_name'])[['mes', 'rule_name', 'sesiones', 'usuarios']]
    df_det.to_csv(os.path.join(out_dir, 'feedback_{}_reglas.csv'.format(sufijo)),
                  index=False, encoding='utf-8-sig')
    escribir_excel(path_xlsx, args.tema, resumen, df_det, df_txt)

    print('')
    print('{:<18}{:>15}{:>17}{:>10}{:>16}'.format('Periodo', 'Total_Usuarios', 'TasaEfectividad', 'Esfuerzo', 'Satisfaccion %'))
    for titulo, m in resumen:
        fmt = lambda v, p: '-' if v is None else (('{:.2%}' if p else '{:.2f}').format(v))
        print('{:<18}{:>15,}{:>17}{:>10}{:>16}'.format(
            titulo, m['Total_Usuarios'], fmt(m['TasaEfectividad'], True),
            fmt(m['Esfuerzo'], False), fmt(m['Satisfaccion'], True)))
    print('')
    if df_txt is not None:
        log('Sugerencias/comentarios: {} filas'.format(len(df_txt)))
    if hay_faltantes:
        log('[ATENCION] Hubo reglas no encontradas: los porcentajes pueden estar mal.')
    log('Excel: {}'.format(path_xlsx))


if __name__ == '__main__':
    main()
